# # Copyright (c) 2026, UpGo Technologies and contributors
# # For license information, please see license.txt

# # import frappe


# def execute(filters=None):
# 	columns, data = [], []
# 	return columns, data



import frappe
from frappe.utils import fmt_money

def execute(filters=None):

    filters = filters or {}

    columns = [
        {"label": "Project", "fieldname": "project_number", "fieldtype": "Link", "options": "FT Project", "width": 200},
        {"label": "Drawing Number", "fieldname": "drawing_number", "width": 200},
        {"label": "PO Serial No", "fieldname": "po_serial_no", "width": 190},
        {"label": "Unit Weight", "fieldname": "unit_weight", "fieldtype": "Float", "width": 200},
        {"label": "Required Qty", "fieldname": "required_qty", "fieldtype": "Float", "width": 190},
        {"label": "Total Weight", "fieldname": "total_weight", "fieldtype": "Float", "width": 200},
    ]

    conditions = ""
    values = {}

    if filters.get("project_number"):
        conditions += " AND p.name IN %(project_number)s"
        values["project_number"] = tuple(filters.get("project_number"))

    if filters.get("drawing_number"):
        conditions += " AND pod.drawing_number IN %(drawing_number)s"
        values["drawing_number"] = tuple(filters.get("drawing_number"))

    if filters.get("is_active"):
        conditions += " AND p.is_active = 1"

    data = frappe.db.sql(f"""
        SELECT
            p.name as project_number,
            ad.drawing_number,
            pod.po_serial_no,
            IFNULL(pod.unit_weight,0) as unit_weight,
            IFNULL(pod.required_qty,0) as required_qty,
            IFNULL(pod.total_weight,0) as total_weight
        FROM `tabFT Project` p
        LEFT JOIN `tabFT Po Drawing` pod
            ON p.name = pod.project_number
        LEFT JOIN `tabFT Add Drawing` ad
            ON ad.name = pod.drawing_number
        WHERE 1=1 {conditions}
        ORDER BY p.name, CAST(pod.po_serial_no AS UNSIGNED)
    """, values, as_dict=True)

    total_unit_weight = sum(d.get("unit_weight", 0) for d in data)
    total_qty         = sum(d.get("required_qty", 0) for d in data)
    total_weight      = sum(d.get("total_weight", 0) for d in data)

    data.append({
        "project_number": "TOTAL",
        "drawing_number": "",
        "po_serial_no":   "",
        "unit_weight":    total_unit_weight,
        "required_qty":   total_qty,
        "total_weight":   total_weight
    })

    return columns, data, None, None, None


@frappe.whitelist()
def download_drawing_excel(filters):

	import frappe
	import openpyxl
	from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
	from io import BytesIO

	filters = frappe.parse_json(filters)

	conditions = ""
	values = {}

	if filters.get("project_number"):
		conditions += " AND p.name IN %(project_number)s"
		values["project_number"] = tuple(filters.get("project_number"))

	if filters.get("drawing_number"):
		conditions += " AND pod.drawing_number IN %(drawing_number)s"
		values["drawing_number"] = tuple(filters.get("drawing_number"))

	data = frappe.db.sql(f"""
		SELECT
			p.name as project,
			IFNULL(pod.drawing_number,'') as drawing_number,
			IFNULL(pod.po_serial_no,'') as po_serial_no,
			IFNULL(pod.unit_weight,0) as unit_weight,
			IFNULL(pod.required_qty,0) as required_qty,
			IFNULL(pod.total_weight,0) as total_weight
		FROM `tabFT Project` p
		LEFT JOIN `tabFT Po Drawing` pod 
			ON p.name = pod.project_number
		WHERE 1=1 {conditions}
		ORDER BY p.name
	""", values, as_dict=True)

	# ---------------- WORKBOOK ----------------
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "PO Drawing Data"

	# ---------------- HEADING ----------------
	ws.merge_cells('A1:F1')

	title_cell = ws['A1']
	title_cell.value = "PO Drawing"
	title_cell.font = Font(size=16, bold=True)
	title_cell.alignment = Alignment(horizontal="center", vertical="center")

	# ---------------- HEADERS ----------------
	headers = [
		"Project",
		"Drawing Number",
		"PO Serial No",
		"Unit Weight",
		"Required Qty",
		"Total Weight"
	]

	header_font = Font(bold=True, size=13, color="FFFFFF")
	header_fill = PatternFill(start_color="2F75B5", end_color="2F75B5", fill_type="solid")

	data_font = Font(size=12)
	total_font = Font(bold=True, size=13, color="FFFFFF")
	total_fill = PatternFill(start_color="2F75B5", end_color="2F75B5", fill_type="solid")

	border = Border(
		left=Side(style="thin"),
		right=Side(style="thin"),
		top=Side(style="thin"),
		bottom=Side(style="thin")
	)

	center = Alignment(horizontal="center", vertical="center")

	# Header Row
	for col, header in enumerate(headers, 1):
		cell = ws.cell(row=2, column=col, value=header)
		cell.font = header_font
		cell.fill = header_fill
		cell.border = border
		cell.alignment = center

	row_no = 3

	unit_weight_total = 0
	required_qty_total = 0
	grand_total = 0

	# ---------------- DATA ----------------
	for d in data:

		unit_weight = d.get("unit_weight") or 0
		required_qty = d.get("required_qty") or 0
		total_weight = d.get("total_weight") or 0

		unit_weight_total += unit_weight
		required_qty_total += required_qty
		grand_total += total_weight

		row_values = [
			d.get("project"),
			d.get("drawing_number"),
			d.get("po_serial_no"),
			unit_weight,
			required_qty,
			total_weight
		]

		for col, val in enumerate(row_values, 1):

			cell = ws.cell(row=row_no, column=col, value=val)
			cell.font = data_font
			cell.border = border
			cell.alignment = center

			if col in [4,6]:
				cell.number_format = '0.000'

		row_no += 1

	# ---------------- TOTAL ROW ----------------
	total_label = ws.cell(row=row_no, column=1, value="Total")
	total_label.font = total_font
	total_label.fill = total_fill
	total_label.alignment = Alignment(horizontal="left", vertical="center")

	ws.cell(row=row_no, column=2).fill = total_fill
	ws.cell(row=row_no, column=3).fill = total_fill

	unit_total_cell = ws.cell(row=row_no, column=4, value=unit_weight_total)
	unit_total_cell.font = total_font
	unit_total_cell.fill = total_fill
	unit_total_cell.number_format = '0.000'
	unit_total_cell.alignment = center

	qty_total_cell = ws.cell(row=row_no, column=5, value=required_qty_total)
	qty_total_cell.font = total_font
	qty_total_cell.fill = total_fill
	qty_total_cell.alignment = center

	total_cell = ws.cell(row=row_no, column=6, value=grand_total)
	total_cell.font = total_font
	total_cell.fill = total_fill
	total_cell.number_format = '0.000'
	total_cell.alignment = center

	for col in range(1, 7):
		ws.cell(row=row_no, column=col).border = border

	# ---------------- COLUMN WIDTH ----------------
	ws.column_dimensions["A"].width = 20
	ws.column_dimensions["B"].width = 22
	ws.column_dimensions["C"].width = 18
	ws.column_dimensions["D"].width = 15
	ws.column_dimensions["E"].width = 15
	ws.column_dimensions["F"].width = 18

	# ---------------- FILE EXPORT ----------------
	file_stream = BytesIO()
	wb.save(file_stream)
	file_stream.seek(0)

	frappe.response['filename'] = "Drawing_Report.xlsx"
	frappe.response['filecontent'] = file_stream.getvalue()
	frappe.response['type'] = 'binary'

   