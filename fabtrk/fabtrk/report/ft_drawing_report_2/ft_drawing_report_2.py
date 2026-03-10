import frappe
from frappe.utils import fmt_money

def execute(filters=None):
    filters = filters or {}

    columns = [
        {"label": "Project", "fieldname": "project_name", "fieldtype": "Link", "options": "FT Project", "width": 200},
        {"label": "Drawing Number", "fieldname": "drawing_number", "fieldtype": "Data","width": 200,"align": "center"},
        {"label": "PO Serial No", "fieldname": "po_serial_no", "fieldtype": "Data","width": 200,"align": "center"},
        {"label": "Unit Weight", "fieldname": "unit_weight", "fieldtype": "Float", "width": 200,"align": "center"},
        {"label": "Required Qty", "fieldname": "quantity", "fieldtype": "Int", "width": 200,"align": "center"},
        {"label": "Total Weight", "fieldname": "total_weight", "fieldtype": "Float", "width": 200,"align": "center"},
    ]

    conditions = ""
    values = {}

    if filters.get("project_number"):
        conditions += " AND p.name IN %(project_number)s"
        values["project_number"] = tuple(filters.get("project_number"))

    if filters.get("drawing_number"):
        conditions += " AND ad.drawing_number IN %(drawing_number)s"
        values["drawing_number"] = tuple(filters.get("drawing_number"))

    if filters.get("is_active"):
        conditions += " AND p.is_active = 1"

    query = f"""
        SELECT
            p.name AS project_name,
            ad.name as drawing_id,
            ad.drawing_number AS drawing_number,
            pod.po_serial_no AS po_serial_no,
            IFNULL(pod.unit_weight, 0) AS unit_weight,
            IFNULL(pod.required_qty, 0) AS quantity,
            (IFNULL(pod.unit_weight,0) * IFNULL(pod.required_qty,0)) AS total_weight
        FROM `tabFT Project` p
        LEFT JOIN `tabFT Add Drawing` ad
            ON ad.project_number = p.name
        LEFT JOIN `tabFT Po Drawing` pod
            ON pod.drawing_number = ad.name
        WHERE 1=1 {conditions}
        ORDER BY CAST(pod.po_serial_no AS UNSIGNED) ASC
    """

    raw_data = frappe.db.sql(query, values, as_dict=1) or []
        # ===== TOTAL Row (Unit Weight + Total Weight) =====
    total_unit_weight = sum(d.get("unit_weight", 0) for d in raw_data)
    total_total_weight = sum(d.get("total_weight", 0) for d in raw_data)

    raw_data.append({
        "project_name": "TOTAL",
        "drawing_number": "",
        "po_serial_no": "",
        "unit_weight": total_unit_weight,
        "quantity": None,
        "total_weight": total_total_weight
    })

    

# ================= SUMMARY =================

    project_names = {d.get("project_name") for d in raw_data if d.get("project_name") != "TOTAL"}
    drawing_ids = {d.get("drawing_id") for d in raw_data if d.get("drawing_id")}
    po_rows = [d for d in raw_data if d.get("po_serial_no")]

    total_projects = len(project_names)
    total_no_of_drawings = len(drawing_ids)
    total_po_drawings = len(po_rows)

    # Total Weight as per Project
    project_conditions = ""
    project_values = {}
    if filters.get("project_number"):
        project_conditions += " AND name IN %(project_number)s"
        project_values["project_number"] = tuple(filters.get("project_number"))
    if filters.get("is_active"):
        project_conditions += " AND is_active = 1"

    project_total_weight = frappe.db.sql(
        f"SELECT SUM(total_weight) FROM `tabFT Project` WHERE 1=1 {project_conditions}",
        project_values,
    )
    project_total_weight = project_total_weight[0][0] if project_total_weight and project_total_weight[0][0] else 0
    
    total_weight_po_items = sum(d.get("total_weight", 0) for d in po_rows)

    
    # Drawing balance for customer = project total - PO items total
    drawing_balance_for_customer = project_total_weight - total_weight_po_items

    formatted_total_weight_as_per_project = "{:,.3f}".format(project_total_weight)
    formatted_total_weight_po_items = "{:,.3f}".format(total_weight_po_items)
    formatted_drawing_balance_for_customer = "{:,.3f}".format(drawing_balance_for_customer)

    report_summary = [
        {
            "label": "",
            "value": f"""
            <div class="summary-container">
                <div class="summary-section">
                    <div class="section-content-count">
                        <p>Total Projects</p>
                        <span>{total_projects}</span>                  
                    </div>
                    <div class="section-content-count">
                        <p>Total Weight as per Project(Kg)</p>
                        <span>{formatted_total_weight_as_per_project}</span>                   
                    </div>
                </div>

                <div class="summary-section">
                    <div class="section-content-count">
                        <p>Project Total No of Line Items</p>
                        <span>{total_po_drawings}</span>                    
                    </div>
                    <div class="section-content-count">
                        <p>Total Weight of Line Items<br>Assign to Project(Kg)</p>
                        <span>{formatted_total_weight_po_items}</span>                    
                    </div>
                </div>
                
                <div class="summary-section">
                    <div class="section-content-count">
                        <p>Total No of Unique Drawings</p>
                        <span>{total_no_of_drawings}</span>                    
                    </div>
                </div>
                
                <div class="summary-section">
                    <div class="section-content-count">
                        <p>Drawing Balance for Customer(Kg)</p>
                        <span>{formatted_drawing_balance_for_customer}</span>                   
                    </div>
                </div>

            </div>
            """,
            "datatype": "HTML",
        }
    ]
    return columns, raw_data, None, None, report_summary

@frappe.whitelist()
def download_drawing_excel(filters):

	import frappe
	import openpyxl
	from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
	from io import BytesIO

	filters = frappe.parse_json(filters)

	conditions = ""
	values = {}

	if filters.get("project_number"):
		conditions += " AND p.name IN %(project_number)s"
		values["project_number"] = tuple(filters.get("project_number"))

	if filters.get("drawing_number"):
		conditions += " AND pod.drawing_number IN %(drawing_number)s"
		values["drawing_number"] = tuple(filters.get("drawing_number"))

	data = frappe.db.sql(f"""
		SELECT
			p.name as project,
			pod.drawing_number,
			pod.po_serial_no,
			pod.unit_weight,
			pod.required_qty,
			pod.total_weight
		FROM `tabFT Po Drawing` pod
		LEFT JOIN `tabFT Project` p 
			ON p.name = pod.project_number
		WHERE 1=1 {conditions}
		ORDER BY p.name
	""", values, as_dict=True)

	# ---------------- WORKBOOK ----------------
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Drawing Report"

	headers = [
		"Project",
		"Drawing Number",
		"PO Serial No",
		"Unit Weight",
		"Required Qty",
		"Total Weight"
	]

	header_font = Font(bold=True, size=13, color="FFFFFF")
	header_fill = PatternFill(start_color="2F75B5", end_color="2F75B5", fill_type="solid")

	data_font = Font(size=12)
	total_font = Font(bold=True, size=13)
	total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

	border = Border(
		left=Side(style="thin"),
		right=Side(style="thin"),
		top=Side(style="thin"),
		bottom=Side(style="thin")
	)

	center = Alignment(horizontal="center", vertical="center")

	# Header
	for col, header in enumerate(headers, 1):
		cell = ws.cell(row=1, column=col, value=header)
		cell.font = header_font
		cell.fill = header_fill
		cell.border = border
		cell.alignment = center

	# Data
	row_no = 2
	grand_total = 0

	for d in data:

		total_weight = d.get("total_weight") or 0
		grand_total += total_weight

		row_values = [
			d.get("project"),
			d.get("drawing_number"),
			d.get("po_serial_no"),
			d.get("unit_weight"),
			d.get("required_qty"),
			total_weight
		]

		for col, val in enumerate(row_values, 1):
			cell = ws.cell(row=row_no, column=col, value=val)
			cell.font = data_font
			cell.border = border
			cell.alignment = center
		row_no += 1

	# Total Row
	ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=5)

	total_label = ws.cell(row=row_no, column=1, value="Total")
	total_label.font = total_font
	total_label.fill = total_fill
	total_label.alignment = start = Alignment(horizontal="left", vertical="center")

	total_cell = ws.cell(row=row_no, column=6, value=grand_total)
	total_cell.font = total_font
	total_cell.fill = total_fill
	total_cell.alignment = center

	for col in range(1, 7):
		ws.cell(row=row_no, column=col).border = border

	# Column Width
	ws.column_dimensions["A"].width = 20
	ws.column_dimensions["B"].width = 20
	ws.column_dimensions["C"].width = 18
	ws.column_dimensions["D"].width = 15
	ws.column_dimensions["E"].width = 15
	ws.column_dimensions["F"].width = 18

	file_stream = BytesIO()
	wb.save(file_stream)
	file_stream.seek(0)

	frappe.response['filename'] = "Drawing_Report.xlsx"
	frappe.response['filecontent'] = file_stream.getvalue()
	frappe.response['type'] = 'binary'