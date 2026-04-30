# Copyright (c) 2026, UpGo Technologies and contributors
# For license information, please see license.txt

# import frappe
# from frappe.model.document import Document
# from frappe.utils.xlsxutils import make_xlsx

# class FTStockRMList(Document):
# 	pass


import frappe
import openpyxl
import csv
import os
import json
from frappe.model.document import Document


class FTStockRMList(Document):
    pass


# EXPORT
@frappe.whitelist()
def export_with_value(file_type="xlsx", export_type="all", filters=None, selected_fields=None):
    import io
    import csv as csv_module
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    meta = frappe.get_meta("FT Stock RM List")

    if selected_fields:
        sf_list = json.loads(selected_fields) if isinstance(selected_fields, str) else selected_fields
    else:
        sf_list = [f.fieldname for f in meta.fields]

    field_label_map = {}
    for f in meta.fields:
        field_label_map[f.fieldname] = (f.label or f.fieldname)
    field_label_map["name"] = "ID"

    valid_fieldnames = ["name"] + [f.fieldname for f in meta.fields]
    fieldnames = [fn for fn in sf_list if fn in valid_fieldnames]

    if not fieldnames:
        frappe.throw("Koi field select nahi ki export ke liye.")

    headers_original = [field_label_map.get(fn, fn) for fn in fieldnames]
    headers_upper = [h.upper() for h in headers_original]

    # ── Records fetch ──
    if export_type == "blank":
        records = []
    else:
        query_kwargs = {"fields": fieldnames}

        if export_type == "filtered" and filters:
            if isinstance(filters, str):
                parsed = json.loads(filters)
            else:
                parsed = filters

            from collections import defaultdict
            equals_groups = defaultdict(list)
            other_filters = []

            for f in parsed:
                if not isinstance(f, list):
                    continue
                if len(f) == 4:
                    _, field, op, value = f
                elif len(f) == 3:
                    field, op, value = f
                else:
                    continue

                if op in ('=', 'Equals', 'equals'):
                    equals_groups[field].append(value)
                else:
                    other_filters.append([field, op, value])

            clean_filters = []
            for field, values in equals_groups.items():
                if len(values) == 1:
                    clean_filters.append([field, '=', values[0]])
                else:
                    clean_filters.append([field, 'in', values])

            for f in other_filters:
                clean_filters.append(f)

            query_kwargs["filters"] = clean_filters if clean_filters else []

        elif export_type == "5":
            query_kwargs["limit"] = 5

        records = frappe.get_all("FT Stock RM List", **query_kwargs)

    # ── Build data rows ──
    data = []
    for d in records:
        row = []
        for field in fieldnames:
            value = d.get(field)
            if field == "grade" and value:
                value = frappe.db.get_value(
                    "FT Material Grade Catalogues", value, "grade"
                ) or value
            if field == "thk":
                value = "THK" if value in [1, True, "1"] else ""
            row.append(value if value is not None else "")
        data.append(row)

    # ── CSV Export — original case headers sab ke liye ──
    if file_type == "csv":
        output = io.StringIO()
        writer = csv_module.writer(output)
        writer.writerow(headers_original)  # ← ALWAYS original case
        writer.writerows(data)
        frappe.local.response.filename    = "FT_Stock_RM_List.csv"
        frappe.local.response.filecontent = output.getvalue().encode("utf-8-sig")
        frappe.local.response.type        = "binary"
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "FT Stock RM List"

    # ── BLANK TEMPLATE ──
    if export_type == "blank":
        ws.append(headers_original)
        for col_idx in range(1, len(headers_original) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        frappe.local.response.filename    = "FT_Stock_RM_List_Template.xlsx"
        frappe.local.response.filecontent = output.getvalue()
        frappe.local.response.type        = "binary"
        return

    # ── STYLED EXCEL ──
    header_fill  = PatternFill("solid", fgColor="BDD7EE")
    header_font  = Font(bold=True, size=10, color="000000")
    data_font    = Font(size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin         = Side(style="thin", color="000000")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(headers_upper)
    for cell in ws[1]:
        cell.fill      = header_fill
        cell.font      = header_font
        cell.alignment = center_align
        cell.border    = border

    ws.row_dimensions[1].height = 30

    if export_type == "all":
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers_upper))}1"

    ws.freeze_panes = "A2"

    for row_idx, row in enumerate(data, start=2):
        ws.append(row)
        fill_color = (
            PatternFill("solid", fgColor="FFFFFF")
            if row_idx % 2 == 0
            else PatternFill("solid", fgColor="F2F2F2")
        )
        for cell in ws[row_idx]:
            cell.fill      = fill_color
            cell.font      = data_font
            cell.border    = border
            cell.alignment = (
                center_align if isinstance(cell.value, (int, float)) else left_align
            )

    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col_cells:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 4, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    frappe.local.response.filename    = "FT_Stock_RM_List.xlsx"
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type        = "binary"


@frappe.whitelist()
def get_field_distinct_values(fieldname, search_text=None):
    meta = frappe.get_meta("FT Stock RM List")
    valid_fields = [f.fieldname for f in meta.fields] + ["name"]

    if fieldname not in valid_fields:
        frappe.throw("Invalid fieldname")

    where_clauses = [f"`{fieldname}` IS NOT NULL", f"`{fieldname}` != ''"]
    params = {}

    if search_text:
        where_clauses.append(f"`{fieldname}` LIKE %(s)s")
        params["s"] = f"%{search_text}%"

    sql = f"""
        SELECT DISTINCT `{fieldname}`
        FROM `tabFT Stock RM List`
        WHERE {' AND '.join(where_clauses)}
        ORDER BY `{fieldname}` ASC
    """

    rows = frappe.db.sql(sql, params)
    return [str(r[0]) for r in rows if r[0] is not None]


# GET FILE HEADERS
@frappe.whitelist()
def get_file_headers(file_url):
    try:
        file_doc  = frappe.get_doc("File", {"file_url": file_url})
        file_path = file_doc.get_full_path()
    except Exception as e:
        frappe.throw(f"File not found: {str(e)}")

    file_ext = os.path.splitext(file_url)[1].lower()
    headers  = []

    if file_ext == ".xlsx":
        wb      = openpyxl.load_workbook(file_path)
        ws      = wb.active
        headers = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(values_only=True))]
    elif file_ext == ".csv":
        with open(file_path, "r", encoding="utf-8-sig") as f:
            headers = [h.strip() for h in next(csv.reader(f))]
    else:
        frappe.throw("Only .xlsx or .csv files are allowed.")

    meta = frappe.get_meta("FT Stock RM List", cached=False)

    label_to_field = {}
    for f in meta.fields:
        if f.label:
            label_to_field[f.label.strip()] = f.fieldname
        label_to_field[f.fieldname.strip()]  = f.fieldname

    all_fields = [
        {"label": f.label, "fieldname": f.fieldname}
        for f in meta.fields
        if f.fieldtype not in ("Section Break", "Column Break", "HTML", "Heading")
    ]

    matched   = {}
    unmatched = []

    for h in headers:
        if not h:
            continue
        if h in label_to_field:
            matched[h] = label_to_field[h]
        else:
            unmatched.append(h)

    return {
        "headers"   : headers,
        "matched"   : matched,
        "unmatched" : unmatched,
        "all_fields": all_fields
    }


# computed_name
def compute_name_python(d, grade_doc=None):
    section_type  = str(d.get("section_type") or "").strip()
    stock_rm_type = str(d.get("stock_rm_type") or "").strip()
    thk_val       = d.get("thk")
    thk_text      = "THK" if thk_val in [1, True, "1", "True"] else ""

    base_name = stock_rm_type
    size_part = ""

    if section_type == "Section":
        name1 = str(d.get("name1") or "").strip()
        if name1:
            size_part = name1 + (" " + thk_text if thk_text else "")

    elif section_type == "Plate":
        thickness = d.get("thickness_mm")
        if thickness is not None and str(thickness).strip() not in ["", "None"]:
            try:
                thickness_str = str(float(thickness))
                if thickness_str.endswith(".0"):
                    thickness_str = thickness_str[:-2]
                if "." in thickness_str:
                    parts         = thickness_str.split(".")
                    thickness_str = parts[0].zfill(2) + "." + parts[1]
                else:
                    thickness_str = thickness_str.zfill(2)
                size_part = thickness_str + (" " + thk_text if thk_text else "")
            except Exception as e:
                frappe.log_error(
                    title="compute_name thickness error",
                    message=f"thickness value: {thickness!r}, error: {e}"
                )

    if not grade_doc:
        return " ".join(filter(None, [base_name, size_part]))

    grade_display = str(grade_doc.get("grade") or "").strip()
    main_bis      = str(grade_doc.get("bis") or "").strip()
    type_bis      = str(
        (grade_doc.get("bis_section") if section_type == "Section"
         else grade_doc.get("bis_plate")) or ""
    ).strip()

    return " ".join(filter(None, [base_name, size_part, type_bis, main_bis, grade_display]))


# IMPORT — Dynamic with mapping support
@frappe.whitelist()
def import_with_value(file_url, custom_mapping=None):

    try:
        file_doc  = frappe.get_doc("File", {"file_url": file_url})
        file_path = file_doc.get_full_path()
    except Exception as e:
        frappe.throw(f"File nahi mili: {str(e)}")

    rows     = []
    file_ext = os.path.splitext(file_url)[1].lower()

    if file_ext == ".xlsx":
        try:
            wb   = openpyxl.load_workbook(file_path)
            ws   = wb.active
            rows = list(ws.iter_rows(values_only=True))
        except Exception as e:
            frappe.throw(f"Excel open nahi hui: {str(e)}")
    elif file_ext == ".csv":
        try:
            with open(file_path, "r", encoding="utf-8-sig") as f:
                rows = list(csv.reader(f))
        except Exception as e:
            frappe.throw(f"CSV open nahi hui: {str(e)}")
    else:
        frappe.throw("Only .xlsx or .csv format ki file upload karo.")

    if not rows:
        frappe.throw("File mein koi data nahi hai.")

    meta = frappe.get_meta("FT Stock RM List", cached=False)

    label_to_field = {}
    for f in meta.fields:
        if f.label:
            label_to_field[f.label.strip()] = f.fieldname
        label_to_field[f.fieldname.strip()] = f.fieldname

    if custom_mapping:
        extra = json.loads(custom_mapping) if isinstance(custom_mapping, str) else custom_mapping
        label_to_field.update(extra)

    float_fields  = ["thickness_mm", "kg__sqm", "kg__meter",
                     "surface_area_sqm__mtr", "surface_area_sqm__ton"]
    check_fields  = ["thk"]
    string_fields = ["name1", "stock_rm_type", "section_type"]

    excel_headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    matched   = [h for h in excel_headers if h and label_to_field.get(h)]
    unmatched = [h for h in excel_headers if h and not label_to_field.get(h)]

    success      = 0
    skipped      = 0
    errors       = []
    skipped_list = []

    # ── Helper: computed_name se BIS section extract karo ──
    def extract_bis_from_computed_name(computed_name_str):
        """
        Format: 'Angle 50*50*6 THK IS808 IS2062 E250A'
        THK ke baad pehla IS-word = bis_section
        """
        if not computed_name_str:
            return None
        parts = str(computed_name_str).strip().split()
        try:
            thk_idx = next(
                (i for i, p in enumerate(parts) if p.upper() == "THK"), None
            )
            if thk_idx is not None and thk_idx + 1 < len(parts):
                candidate = parts[thk_idx + 1]
                if candidate.upper().startswith("IS") or candidate.upper() == "HARDOX":
                    return candidate
        except Exception:
            pass
        return None

    # ── Helper: sahi grade dhundo — bis_section se match karke ──
    def find_grade_id(raw_grade, bis_section_hint=None):
        """
        1. Grade name + bis_section dono se match karo (most accurate)
        2. Grade name se match karo (pehla result)
        3. Case-insensitive fallback
        """
        raw_grade = str(raw_grade).strip()

        # Step 1: grade + bis_section dono match
        if bis_section_hint:
            result = frappe.db.sql("""
                SELECT name FROM `tabFT Material Grade Catalogues`
                WHERE grade = %s
                AND (
                    LOWER(TRIM(bis_section)) = LOWER(TRIM(%s))
                    OR LOWER(TRIM(bis_plate)) = LOWER(TRIM(%s))
                )
                LIMIT 1
            """, (raw_grade, bis_section_hint, bis_section_hint))
            if result:
                return result[0][0]

        # Step 2: sirf grade name se (pehla match)
        grade_id = frappe.db.get_value(
            "FT Material Grade Catalogues",
            {"grade": raw_grade},
            "name"
        )
        if grade_id:
            return grade_id

        # Step 3: grade name se ID match
        grade_id = frappe.db.get_value(
            "FT Material Grade Catalogues",
            {"name": raw_grade},
            "name"
        )
        if grade_id:
            return grade_id

        # Step 4: case-insensitive fallback
        result = frappe.db.sql("""
            SELECT name FROM `tabFT Material Grade Catalogues`
            WHERE LOWER(TRIM(grade)) = LOWER(TRIM(%s))
            LIMIT 1
        """, (raw_grade,))
        if result:
            return result[0][0]

        return None

    for i, row in enumerate(rows[1:], start=2):
        try:
            d = {}

            if not row or all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            for idx, excel_label in enumerate(excel_headers):
                if not excel_label:
                    continue
                fieldname = label_to_field.get(excel_label.strip())
                if not fieldname:
                    continue

                val = row[idx] if idx < len(row) else None

                if fieldname in float_fields:
                    try:
                        val = float(val) if val not in [None, ""] else None
                    except:
                        val = None

                elif fieldname in check_fields:
                    val_str = str(val).strip().lower() if val not in [None, ""] else ""
                    val = 1 if val_str in ["1", "true", "thk"] else 0

                elif fieldname in string_fields:
                    val = str(val).strip() if val not in [None, ""] else None

                if val == "" or str(val) == "None":
                    val = None

                d[fieldname] = val

            if not any(v is not None for v in d.values()):
                continue

            # ── Grade lookup — bis_section hint use karo ──
            grade_doc     = None
            grade_display = ""

            if d.get("grade"):
                raw_grade = str(d["grade"]).strip()

                # Excel mein computed_name column hai to usse BIS hint lo
                bis_hint = None
                computed_name_col = label_to_field.get("Computed Name") or label_to_field.get("computed_name")
                if computed_name_col:
                    # computed_name ki value row mein dhundo
                    cn_idx = next(
                        (idx for idx, h in enumerate(excel_headers)
                         if label_to_field.get(h) == "computed_name"),
                        None
                    )
                    if cn_idx is not None and cn_idx < len(row):
                        cn_val = row[cn_idx]
                        if cn_val:
                            bis_hint = extract_bis_from_computed_name(str(cn_val))

                grade_id = find_grade_id(raw_grade, bis_hint)

                if grade_id:
                    grade_doc = frappe.db.get_value(
                        "FT Material Grade Catalogues",
                        grade_id,
                        ["grade", "bis", "bis_section", "bis_plate"],
                        as_dict=True
                    )
                    grade_display = grade_doc.get("grade") or raw_grade
                    d["grade"] = grade_id
                else:
                    errors.append(
                        f"⚠️ Row {i}: Grade '{raw_grade}' FT Material Grade Catalogues mein nahi mila — skipped"
                    )
                    continue

            # ── Computed name ──
            d["computed_name"] = compute_name_python(d, grade_doc)

            # ── Duplicate check — computed_name + grade dono ──
            duplicate_filters = {"computed_name": d["computed_name"]}
            if d.get("grade"):
                duplicate_filters["grade"] = d["grade"]

            existing = frappe.db.get_value(
                "FT Stock RM List",
                duplicate_filters,
                "name"
            )
            if existing:
                skipped += 1
                skipped_list.append(
                    f"⚠️ Row {i}: '{d['computed_name']}'"
                    + (f" | Grade: {grade_display}" if grade_display else "")
                    + f" — already exists ({existing})"
                )
                continue

            d.pop("name", None)
            d.pop("naming_series", None)

            doc = frappe.new_doc("FT Stock RM List")
            doc.update(d)
            doc.insert(ignore_permissions=True)
            success += 1

        except Exception as e:
            errors.append(f"❌ Row {i}: {str(e)}")
            frappe.log_error(
                title=f"Import Error Row {i}",
                message=f"Row: {row}\nError: {str(e)}"
            )

    frappe.db.commit()

    # ── Summary ──
    msg = f"✅ {success} records successfully imported."

    if skipped > 0:
        msg += f"\n\n⚠️ {skipped} records already exist (skipped):"
        msg += "\n" + "\n".join(skipped_list[:20])
        if skipped > 20:
            msg += f"\n... aur {skipped - 20} more skipped"

    msg += f"\n\n📋 Matched fields ({len(matched)}): {', '.join(matched)}"

    if unmatched:
        msg += f"\n⚠️ Unmatched fields ({len(unmatched)}): {', '.join(unmatched)}"

    if errors:
        msg += f"\n\n❌ Errors ({len(errors)}):\n" + "\n".join(errors[:15])
        if len(errors) > 15:
            msg += f"\n... aur {len(errors) - 15} errors"

    return msg




    